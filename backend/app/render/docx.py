"""Word rendering.

Word is where teaching material actually gets finished — a teacher will tweak a question,
add their school's header, and print. So the priorities here are different from the slide
renderer: real Word styles (so the document has a navigable outline and the teacher can
restyle it in one click), genuine page breaks, and answer lines that survive editing.

The Markdown in note bodies is converted to Word runs rather than dumped as literal
asterisks — a small thing that makes the difference between output that looks generated
and output that looks written.
"""

from __future__ import annotations

import re
from pathlib import Path

from app.core.capabilities import require
from app.core.logging import get_logger
from app.render.theme import Theme, get_theme
from app.schemas.lesson import LessonPlan
from app.schemas.notes import LectureNotes, NoteSection
from app.schemas.paper import Question, QuestionPaper
from app.schemas.rubric import Rubric, RubricStyle
from app.schemas.study import Worksheet

log = get_logger(__name__)

_INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def _document(theme: Theme):  # noqa: ANN202
    """A Document with the base styles set from the theme."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = theme.body_font
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for level, size in ((1, 20), (2, 15), (3, 13)):
        style = document.styles[f"Heading {level}"]
        style.font.name = theme.title_font
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(theme.title)
        style.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        style.paragraph_format.space_after = Pt(6)

    return document


def _rich_text(paragraph, text: str) -> None:  # noqa: ANN001
    """Write text, honouring **bold**, *italic* and `code` inline markup."""
    for piece in _INLINE_RE.split(text):
        if not piece:
            continue
        run = paragraph.add_run()
        if piece.startswith("**") and piece.endswith("**"):
            run.text = piece[2:-2]
            run.bold = True
        elif piece.startswith("*") and piece.endswith("*"):
            run.text = piece[1:-1]
            run.italic = True
        elif piece.startswith("`") and piece.endswith("`"):
            run.text = piece[1:-1]
            run.font.name = "Consolas"
        else:
            run.text = piece


def _markdown_body(document, markdown: str, theme: Theme) -> None:  # noqa: ANN001
    """Render a block of Markdown prose: paragraphs, bullet and numbered lists."""
    for block in markdown.split("\n"):
        line = block.rstrip()
        if not line.strip():
            continue

        bullet = re.match(r"^\s*[-*+]\s+(.*)$", line)
        numbered = re.match(r"^\s*\d+[.)]\s+(.*)$", line)
        quote = re.match(r"^\s*>\s?(.*)$", line)

        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            _rich_text(paragraph, bullet.group(1))
        elif numbered:
            paragraph = document.add_paragraph(style="List Number")
            _rich_text(paragraph, numbered.group(1))
        elif quote:
            paragraph = document.add_paragraph(style="Intense Quote")
            _rich_text(paragraph, quote.group(1))
        else:
            paragraph = document.add_paragraph()
            _rich_text(paragraph, line)


def _callout(document, label: str, body: str, theme: Theme, colour: str) -> None:  # noqa: ANN001
    """A boxed aside, drawn as a single-cell table so it survives editing."""
    from docx.shared import Pt, RGBColor

    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    cell.text = ""

    heading = cell.paragraphs[0]
    run = heading.add_run(label.upper())
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(colour)

    body_paragraph = cell.add_paragraph()
    _rich_text(body_paragraph, body)
    _shade(cell, theme.surface)
    document.add_paragraph()


def _shade(cell, hex_colour: str) -> None:  # noqa: ANN001
    """Cell background. python-docx has no API for this, so set the XML directly."""
    from docx.oxml.ns import qn
    from docx.oxml.parser import OxmlElement

    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(shading)


def _table_from_rows(document, rows: list[list[str]], theme: Theme) -> None:  # noqa: ANN001
    from docx.shared import Pt

    if not rows:
        return
    width = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"

    for row_index, row in enumerate(rows):
        for column_index in range(width):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            _rich_text(paragraph, row[column_index] if column_index < len(row) else "")
            for run in paragraph.runs:
                run.font.size = Pt(9 if row_index else 9.5)
                if row_index == 0:
                    run.bold = True
        if row_index == 0:
            for column_index in range(width):
                _shade(table.cell(0, column_index), theme.surface)
    document.add_paragraph()


# --------------------------------------------------------------------------- notes


def render_notes(notes: LectureNotes, out_path: Path, *, theme_key: str = "academic") -> Path:
    require("docx")
    theme = get_theme(theme_key)
    document = _document(theme)

    document.add_heading(notes.title, level=0)
    if notes.subtitle:
        document.add_paragraph(notes.subtitle).italic = True

    if notes.overview:
        _markdown_body(document, notes.overview, theme)

    if notes.objectives:
        document.add_heading("Learning objectives", level=1)
        for objective in notes.objectives:
            paragraph = document.add_paragraph(style="List Number")
            _rich_text(paragraph, f"{objective.text} ({objective.bloom.value})")

    if notes.prerequisites:
        document.add_heading("Before you start", level=1)
        for item in notes.prerequisites:
            document.add_paragraph(item, style="List Bullet")

    for section in notes.sections:
        _render_note_section(document, section, theme, level=1)

    if notes.summary:
        document.add_heading("Summary", level=1)
        _markdown_body(document, notes.summary, theme)

    if notes.glossary:
        document.add_heading("Glossary", level=1)
        _table_from_rows(
            document,
            [["Term", "Definition"]] + [[t.term, t.definition] for t in notes.glossary],
            theme,
        )

    if notes.review_questions:
        document.add_heading("Review questions", level=1)
        for question in notes.review_questions:
            document.add_paragraph(question.question, style="List Number")
        if any(q.answer for q in notes.review_questions):
            document.add_page_break()
            document.add_heading("Review answers", level=1)
            for question in notes.review_questions:
                paragraph = document.add_paragraph(style="List Number")
                _rich_text(paragraph, f"**{question.question}** — {question.answer}")

    if notes.further_reading:
        document.add_heading("Further reading", level=1)
        for item in notes.further_reading:
            document.add_paragraph(item, style="List Bullet")

    if notes.citations:
        document.add_heading("Sources", level=1)
        for citation in notes.citations:
            document.add_paragraph(citation.label(), style="List Bullet")

    return _save(document, out_path)


_CALLOUT_COLOURS = {
    "misconception": "B91C1C",
    "warning": "B45309",
    "tip": "047857",
    "exam_tip": "6D28D9",
    "real_world": "0369A1",
    "note": "374151",
}


def _render_note_section(document, section: NoteSection, theme: Theme, level: int) -> None:  # noqa: ANN001
    document.add_heading(section.heading, level=min(level, 4))
    if section.body:
        _markdown_body(document, section.body, theme)

    if section.key_terms:
        _table_from_rows(
            document,
            [["Term", "Definition"]] + [[t.term, t.definition] for t in section.key_terms],
            theme,
        )

    for example in section.examples:
        _callout(
            document,
            "Worked example",
            example.prompt
            + "".join(f"\n{i}. {s}" for i, s in enumerate(example.steps, start=1))
            + (f"\nAnswer: {example.answer}" if example.answer else ""),
            theme,
            theme.accent,
        )

    for callout in section.callouts:
        _callout(
            document,
            callout.title or callout.kind.value.replace("_", " "),
            callout.body,
            theme,
            _CALLOUT_COLOURS.get(callout.kind.value, theme.accent),
        )

    for subsection in section.subsections:
        _render_note_section(document, subsection, theme, level + 1)


# --------------------------------------------------------------------------- exams


def render_paper(
    paper: QuestionPaper,
    out_path: Path,
    *,
    answer_key: bool = False,
    theme_key: str = "academic",
) -> Path:
    """Render the paper as printed, or the answer key."""
    require("docx")
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    theme = get_theme(theme_key)
    document = _document(theme)
    meta = paper.meta

    if meta.institution:
        heading = document.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = heading.add_run(meta.institution.upper())
        run.bold = True
        run.font.size = Pt(13)

    title = meta.exam_name or "Examination"
    if answer_key:
        title += " — Answer Key"
    if meta.variant_label:
        title += f"  ({meta.variant_label})"
    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(16)

    header_rows = [
        [
            f"Course: {meta.course}" if meta.course else "",
            f"Class: {meta.grade_level}" if meta.grade_level else "",
        ],
        [
            f"Time: {meta.duration_minutes} minutes" if meta.duration_minutes else "",
            f"Maximum marks: {paper.total_marks:g}",
        ],
    ]
    if meta.date:
        header_rows.append([f"Date: {meta.date}", ""])
    _table_from_rows(document, header_rows, theme)

    if not answer_key:
        name_line = document.add_paragraph()
        name_line.add_run("Name: ").bold = True
        name_line.add_run("_" * 34 + "     ")
        name_line.add_run("Roll No: ").bold = True
        name_line.add_run("_" * 18)

        if meta.instructions:
            document.add_heading("General instructions", level=2)
            for instruction in meta.instructions:
                document.add_paragraph(instruction, style="List Number")

    number = 0
    for section_index, section in enumerate(paper.sections):
        if section_index:
            document.add_paragraph()
        document.add_heading(section.title, level=1)
        if section.instructions and not answer_key:
            document.add_paragraph(section.instructions).italic = True
        if section.choose_count:
            note = document.add_paragraph()
            _rich_text(note, f"*Answer any {section.choose_count} questions from this section.*")

        for question in section.questions:
            number += 1
            _render_question(document, question, number, theme, answer_key=answer_key)

    if not answer_key:
        closing = document.add_paragraph()
        closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
        closing.add_run("— End of paper —").italic = True

    return _save(document, out_path)


def _render_question(  # noqa: PLR0913
    document,  # noqa: ANN001
    question: Question,
    number: int,
    theme: Theme,
    *,
    answer_key: bool,
    answer_lines: int = 0,
) -> None:
    from docx.shared import Pt, RGBColor

    if question.scenario:
        _callout(document, "Read the following", question.scenario, theme, theme.accent)

    stem = document.add_paragraph()
    stem.add_run(f"{number}.  ").bold = True
    _rich_text(stem, question.text)
    marks = stem.add_run(f"   [{question.marks:g}]")
    marks.bold = True
    marks.font.color.rgb = RGBColor.from_string(theme.muted)

    if question.options:
        for index, option in enumerate(question.options):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Pt(28)
            paragraph.paragraph_format.space_after = Pt(2)
            label = paragraph.add_run(f"{chr(ord('A') + index)}.  ")
            label.bold = True
            _rich_text(paragraph, option.text)
            if answer_key and option.is_correct:
                tick = paragraph.add_run("   ✓")
                tick.bold = True
                tick.font.color.rgb = RGBColor.from_string("047857")

    if question.pairs:
        _table_from_rows(
            document,
            [["", "Column A", "", "Column B"]]
            + [
                [f"{i + 1}.", pair.left, f"{chr(ord('a') + i)}.", pair.right]
                for i, pair in enumerate(question.pairs)
            ],
            theme,
        )

    for sub_index, sub in enumerate(question.sub_questions):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(28)
        paragraph.add_run(f"({sub.label or chr(ord('a') + sub_index)})  ").bold = True
        _rich_text(paragraph, sub.text)
        paragraph.add_run(f"   [{sub.marks:g}]").bold = True
        if answer_key and sub.answer:
            answer = document.add_paragraph()
            answer.paragraph_format.left_indent = Pt(44)
            _rich_text(answer, f"*{sub.answer}*")

    if answer_key:
        _answer_block(document, question, theme)
    elif answer_lines:
        for _ in range(answer_lines):
            line = document.add_paragraph("_" * 88)
            line.paragraph_format.space_after = Pt(4)
            line.paragraph_format.left_indent = Pt(28)


def _answer_block(document, question: Question, theme: Theme) -> None:  # noqa: ANN001
    from docx.shared import Pt, RGBColor

    answer = document.add_paragraph()
    answer.paragraph_format.left_indent = Pt(28)
    label = answer.add_run("Answer:  ")
    label.bold = True
    label.font.color.rgb = RGBColor.from_string("047857")
    _rich_text(answer, question.answer_text())

    if question.working:
        working = document.add_paragraph()
        working.paragraph_format.left_indent = Pt(28)
        working.add_run("Working: ").bold = True
        _rich_text(working, question.working.replace("\n", "  →  "))

    if question.rubric_points:
        scheme = document.add_paragraph()
        scheme.paragraph_format.left_indent = Pt(28)
        scheme.add_run("Mark scheme:").bold = True
        for point in question.rubric_points:
            item = document.add_paragraph(style="List Bullet")
            item.paragraph_format.left_indent = Pt(44)
            _rich_text(item, f"{point.description}  *({point.marks:g})*")

    if question.explanation:
        explanation = document.add_paragraph()
        explanation.paragraph_format.left_indent = Pt(28)
        _rich_text(explanation, f"*{question.explanation}*")

    wrong = [o for o in question.options if o.rationale and not o.is_correct]
    if wrong:
        header = document.add_paragraph()
        header.paragraph_format.left_indent = Pt(28)
        header.add_run("Distractor analysis:").bold = True
        for option in wrong:
            item = document.add_paragraph(style="List Bullet")
            item.paragraph_format.left_indent = Pt(44)
            _rich_text(item, f"**{option.text}** — {option.rationale}")

    tags = document.add_paragraph()
    tags.paragraph_format.left_indent = Pt(28)
    run = tags.add_run(f"{question.topic} · {question.bloom.value} · {question.difficulty.value}")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(theme.muted)


# --------------------------------------------------------------------------- worksheet


def render_worksheet(
    worksheet: Worksheet,
    out_path: Path,
    *,
    answer_key: bool = False,
    theme_key: str = "academic",
) -> Path:
    require("docx")
    from docx.shared import Pt

    theme = get_theme(theme_key)
    document = _document(theme)

    document.add_heading(worksheet.title + (" — Answers" if answer_key else ""), level=0)
    subtitle = " · ".join(
        p
        for p in (worksheet.subject, worksheet.grade_level, f"{worksheet.estimated_minutes} min")
        if p
    )
    if subtitle:
        document.add_paragraph(subtitle).italic = True

    if not answer_key:
        name = document.add_paragraph()
        name.add_run("Name: ").bold = True
        name.add_run("_" * 34 + "     ")
        name.add_run("Date: ").bold = True
        name.add_run("_" * 18)

    if worksheet.objectives:
        _callout(
            document,
            "In this worksheet you will",
            "\n".join(f"• {o.text}" for o in worksheet.objectives),
            theme,
            theme.accent,
        )

    if worksheet.warm_up and not answer_key:
        document.add_heading("Warm up", level=1)
        for item in worksheet.warm_up:
            document.add_paragraph(item, style="List Number")

    if worksheet.worked_example:
        _callout(document, "Worked example", worksheet.worked_example, theme, theme.accent_alt)

    number = 0
    for section in worksheet.sections:
        document.add_heading(section.title, level=1)
        if section.instructions:
            document.add_paragraph(section.instructions).italic = True
        for question in section.questions:
            number += 1
            _render_question(
                document,
                question,
                number,
                theme,
                answer_key=answer_key,
                answer_lines=0 if answer_key else section.answer_lines,
            )

    if worksheet.challenge:
        document.add_heading("Challenge", level=1)
        for question in worksheet.challenge:
            number += 1
            _render_question(document, question, number, theme, answer_key=answer_key)

    if worksheet.reflection_prompt and not answer_key:
        document.add_heading("Reflect", level=1)
        document.add_paragraph(worksheet.reflection_prompt)
        for _ in range(3):
            document.add_paragraph("_" * 88).paragraph_format.space_after = Pt(4)

    return _save(document, out_path)


# --------------------------------------------------------------------------- rubric


def render_rubric(rubric: Rubric, out_path: Path, *, theme_key: str = "academic") -> Path:
    require("docx")
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches

    theme = get_theme(theme_key)
    document = _document(theme)

    # A criteria × levels grid needs the width; portrait crushes the descriptors.
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = section.right_margin = Inches(0.6)

    document.add_heading(rubric.title, level=0)
    if rubric.task_description:
        document.add_paragraph(rubric.task_description)
    document.add_paragraph(f"Total: {rubric.total_points:g} points").bold = True

    if rubric.style == RubricStyle.HOLISTIC:
        _table_from_rows(
            document,
            [["Level", "What this looks like"]]
            + [[c.level_name, c.descriptor] for c in rubric.holistic_descriptors],
            theme,
        )
    else:
        _table_from_rows(document, rubric.grid(), theme)

    if rubric.student_facing_summary:
        document.add_heading("What this means for you", level=1)
        _markdown_body(document, rubric.student_facing_summary, theme)

    return _save(document, out_path)


# --------------------------------------------------------------------------- lesson plan


def render_lesson_plan(plan: LessonPlan, out_path: Path, *, theme_key: str = "academic") -> Path:
    require("docx")
    theme = get_theme(theme_key)
    document = _document(theme)

    document.add_heading(plan.title, level=0)
    _table_from_rows(
        document,
        [
            ["Subject", plan.subject, "Class", plan.grade_level],
            ["Duration", f"{plan.duration_minutes} min", "Date", plan.date],
        ],
        theme,
    )

    if plan.objectives:
        document.add_heading("Learning objectives", level=1)
        for objective in plan.objectives:
            paragraph = document.add_paragraph(style="List Bullet")
            _rich_text(paragraph, f"{objective.text} *({objective.bloom.value})*")

    if plan.success_criteria:
        document.add_heading("Success criteria", level=1)
        for criterion in plan.success_criteria:
            document.add_paragraph(criterion, style="List Bullet")

    if plan.materials:
        document.add_heading("Materials", level=1)
        for material in plan.materials:
            document.add_paragraph(material, style="List Bullet")

    if plan.hook:
        _callout(document, "Hook", plan.hook, theme, theme.accent)

    if plan.activities:
        document.add_heading("Lesson sequence", level=1)
        _table_from_rows(
            document,
            [["Time", "Activity", "Teacher does", "Students do", "Check"]]
            + [
                [
                    f"{a.minutes} min",
                    a.name,
                    a.teacher_does,
                    a.students_do,
                    a.check_for_understanding,
                ]
                for a in plan.activities
            ],
            theme,
        )

    if plan.closure:
        _callout(document, "Closure", plan.closure, theme, theme.accent_alt)

    differentiation = plan.differentiation
    groups = [
        ("Support", differentiation.support),
        ("Core", differentiation.core),
        ("Extension", differentiation.extension),
        ("Language support", differentiation.language_support),
        ("Accessibility", differentiation.accessibility),
    ]
    if any(items for _, items in groups):
        document.add_heading("Differentiation", level=1)
        _table_from_rows(
            document,
            [["Group", "Adjustments"]]
            + [[label, "\n".join(f"• {i}" for i in items)] for label, items in groups if items],
            theme,
        )

    if plan.anticipated_misconceptions:
        document.add_heading("Anticipated misconceptions", level=1)
        for item in plan.anticipated_misconceptions:
            document.add_paragraph(item, style="List Bullet")

    if plan.formative_assessment:
        document.add_heading("Assessment", level=1)
        for item in plan.formative_assessment:
            document.add_paragraph(item, style="List Bullet")
    if plan.summative_assessment:
        document.add_paragraph(f"Summative: {plan.summative_assessment}")
    if plan.homework:
        document.add_heading("Homework", level=1)
        document.add_paragraph(plan.homework)
    if plan.teacher_notes:
        document.add_heading("Notes", level=1)
        _markdown_body(document, plan.teacher_notes, theme)

    return _save(document, out_path)


def _save(document, out_path: Path) -> Path:  # noqa: ANN001
    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))
    log.info("Wrote %s", out_path.name)
    return out_path
