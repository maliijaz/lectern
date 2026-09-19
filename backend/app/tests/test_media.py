"""Figure rendering: charts, diagrams, and graceful absence of either."""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import pytest

from app.media import diagrams
from app.schemas.common import Figure
from app.schemas.deck import Deck, Slide, SlideLayout

CHART_SPEC = json.dumps(
    {
        "type": "bar",
        "title": "Rate of photosynthesis by light intensity",
        "x_label": "Light intensity (lux)",
        "y_label": "O2 produced (ml/min)",
        "labels": ["200", "400", "600", "800"],
        "series": [{"name": "Elodea", "values": [2, 5, 8, 9]}],
    }
)


def _matplotlib_available() -> bool:
    import importlib.util

    return importlib.util.find_spec("matplotlib") is not None


@pytest.mark.skipif(not _matplotlib_available(), reason="matplotlib is not installed")
def test_chart_renders_to_png(tmp_path: Path) -> None:
    figure = Figure(kind="chart", spec=CHART_SPEC, alt_text="Bar chart of rate against intensity")
    path = diagrams.render_figure(figure, tmp_path)

    assert path is not None
    assert path.exists()
    assert path.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n"


@pytest.mark.skipif(not _matplotlib_available(), reason="matplotlib is not installed")
def test_chart_is_cached_by_its_source(tmp_path: Path) -> None:
    figure = Figure(kind="chart", spec=CHART_SPEC)
    first = diagrams.render_figure(figure, tmp_path)
    mtime = first.stat().st_mtime_ns

    second = diagrams.render_figure(figure, tmp_path)
    assert second == first
    assert second.stat().st_mtime_ns == mtime, "a cached figure must not be re-rendered"

    # Different source, different file.
    other = diagrams.render_figure(
        Figure(kind="chart", spec=CHART_SPEC.replace("Elodea", "Cabomba")), tmp_path
    )
    assert other != first


def test_empty_spec_renders_nothing(tmp_path: Path) -> None:
    assert diagrams.render_figure(Figure(kind="diagram", spec=""), tmp_path) is None


def test_invalid_chart_spec_is_survivable(tmp_path: Path) -> None:
    """A model that emits broken JSON costs one illustration, not the export."""
    assert diagrams.render_figure(Figure(kind="chart", spec="{not json"), tmp_path) is None


def test_unknown_figure_kind_renders_nothing(tmp_path: Path) -> None:
    assert diagrams.render_figure(Figure(kind="interpretive-dance", spec="x"), tmp_path) is None


@pytest.mark.skipif(shutil.which("mmdc") is not None, reason="mmdc is installed")
def test_mermaid_without_the_cli_degrades_quietly(tmp_path: Path) -> None:
    figure = Figure(kind="diagram", spec="graph TD; A-->B;", alt_text="A leads to B")
    assert diagrams.render_figure(figure, tmp_path) is None


def test_prepare_deck_figures_attaches_paths(tmp_path: Path) -> None:
    """Rendered paths must reach the renderer without entering the stored JSON."""
    deck = Deck(
        title="Test",
        slides=[
            Slide(layout=SlideLayout.BULLETS, heading="No figure"),
            Slide(
                layout=SlideLayout.IMAGE_TEXT,
                heading="With a chart",
                figure=Figure(kind="chart", spec=CHART_SPEC),
            ),
        ],
    )

    count = diagrams.prepare_deck_figures(deck, tmp_path)

    if _matplotlib_available():
        assert count == 1
        attached = getattr(deck.slides[1].figure, "_rendered_path", None)
        assert attached and Path(attached).exists()
    else:
        assert count == 0

    # The rendered path is a runtime detail; it must not be serialised into the artifact.
    assert "_rendered_path" not in deck.model_dump_json()


@pytest.mark.skipif(not _matplotlib_available(), reason="matplotlib is not installed")
def test_deck_with_a_chart_exports_to_pptx(tmp_path: Path) -> None:
    """The whole path: spec → image → placed on the slide."""
    from app.render import pptx

    deck = Deck(
        title="Photosynthesis",
        slides=[
            Slide(
                layout=SlideLayout.IMAGE_TEXT,
                heading="Rate against intensity",
                bullets=["Rate rises then plateaus"],
                figure=Figure(kind="chart", spec=CHART_SPEC, caption="Elodea, 25°C"),
            )
        ],
    )
    diagrams.prepare_deck_figures(deck, tmp_path)
    out = pptx.render(deck, tmp_path / "deck.pptx")

    from pptx import Presentation

    presentation = Presentation(str(out))
    shapes = presentation.slides[0].shapes
    assert any(shape.shape_type == 13 for shape in shapes), "the chart image was not placed"
